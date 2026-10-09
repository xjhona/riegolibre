"""Lista de materiales (metrado) y costos."""

import csv
import math
from dataclasses import dataclass
from typing import Dict, List, Optional

from .precios import (Articulo, candidatos_conector, candidatos_tuberia,
                      precio_representativo)

ORDEN_CATEGORIAS = ["Laterales", "Emisores", "Portalaterales", "Red principal", "Accesorios",
                    "Válvulas y equipos"]


@dataclass
class Partida:
    categoria: str
    material: str  # nombre en el diseño
    cantidad: float  # cantidad neta del diseño
    unidad: str  # "m" o "und"
    clave: str  # identifica la partida para recordar el artículo elegido
    articulo: Optional[Articulo] = None
    automatico: bool = False  # True si el artículo se asignó solo (por diámetro y clase)
    nota: str = ""

    def cantidad_compra(self, desperdicio=0.0):
        """Cantidad a comprar: los metros llevan el porcentaje de desperdicio."""
        return self.cantidad * (1 + desperdicio) if self.unidad == "m" else self.cantidad

    def piezas(self, desperdicio=0.0):
        """Número de tubos, si el artículo se vende por tubo."""
        if self.articulo is None or not self.articulo.por_pieza or self.unidad != "m":
            return None
        return math.ceil(self.cantidad_compra(desperdicio) / self.articulo.longitud_pieza_m - 1e-9)

    def precio_unitario(self):
        """Precio por la unidad de la partida (por metro en tuberías)."""
        if self.articulo is None:
            return None
        return self.articulo.precio_m if self.unidad == "m" else self.articulo.precio

    def costo(self, desperdicio=0.0):
        if self.articulo is None:
            return None
        piezas = self.piezas(desperdicio)
        if piezas is not None:
            return piezas * self.articulo.precio
        return self.cantidad_compra(desperdicio) * self.precio_unitario()


def total(partidas, desperdicio=0.0):
    return sum(p.costo(desperdicio) or 0.0 for p in partidas)


def clave_articulo(articulo):
    """Identificador de un artículo (hay listas con códigos repetidos, como «S.C»)."""
    return f"{articulo.codigo}|{articulo.descripcion}"


def asignar_articulos(partidas, articulos, tuberias_por_nombre, elegidos: Dict[str, str]):
    """Asigna a cada partida su artículo de la lista de precios.

    Primero se respeta la elección guardada (elegidos: clave de partida -> clave de
    artículo). Si no la hay, las tuberías del catálogo se buscan en la lista por
    material, diámetro y clase, y los conectores iniciales por el diámetro de la
    manguera; entre varios equivalentes se toma el de precio central.
    """
    por_clave = {clave_articulo(a): a for a in articulos}
    for partida in partidas:
        partida.articulo, partida.automatico = None, False
        elegido = elegidos.get(partida.clave)
        if elegido is not None:
            partida.articulo = por_clave.get(elegido)
            if partida.articulo is not None:
                continue
        tipo, _, nombre = partida.clave.partition(":")
        tuberia = tuberias_por_nombre.get(nombre)
        if tuberia is None:
            continue
        if tipo == "tuberia":
            candidatos = candidatos_tuberia(tuberia, articulos)
        elif tipo == "conector" and tuberia.diametro_nominal_mm:
            candidatos = candidatos_conector(tuberia.diametro_nominal_mm, articulos)
        else:
            continue
        partida.articulo = precio_representativo(candidatos)
        partida.automatico = partida.articulo is not None


def ordenar(partidas):
    def clave(p):
        indice = ORDEN_CATEGORIAS.index(p.categoria) if p.categoria in ORDEN_CATEGORIAS else 99
        return indice, p.material
    return sorted(partidas, key=clave)


# --------------------------------------------------------------- exportación


COLUMNAS = ["Categoría", "Material", "Cantidad", "Unidad", "Cantidad a comprar", "Tubos",
            "Código", "Artículo de la lista de precios", "Precio unitario (por tubo, m o und)",
            "Subtotal"]


def _fila(partida, desperdicio):
    a = partida.articulo
    piezas = partida.piezas(desperdicio)
    unitario = a.precio if piezas is not None else partida.precio_unitario()
    return [partida.categoria, partida.material, round(partida.cantidad, 2), partida.unidad,
            round(partida.cantidad_compra(desperdicio), 2), piezas,
            a.codigo if a else "", a.descripcion if a else "",
            round(unitario, 4) if unitario is not None else None,
            round(partida.costo(desperdicio), 2) if a else None]


def exportar_csv(partidas, ruta, desperdicio=0.0, moneda=""):
    with open(ruta, "w", encoding="utf-8-sig", newline="") as archivo:
        escritor = csv.writer(archivo, delimiter=";")
        escritor.writerow(COLUMNAS)
        for partida in partidas:
            escritor.writerow(["" if v is None else v for v in _fila(partida, desperdicio)])
        escritor.writerow([""] * 8 + [f"Total {moneda}".strip(), round(total(partidas, desperdicio), 2)])


def exportar_excel(partidas, ruta, desperdicio=0.0, moneda="", titulo="Lista de materiales",
                   detalles=()):
    """Exporta a .xlsx (requiere openpyxl, incluido en QGIS)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Materiales"
    hoja.append([titulo])
    hoja["A1"].font = Font(bold=True, size=14)
    for detalle in detalles:
        hoja.append([detalle])
    hoja.append([f"Desperdicio aplicado a tuberías: {desperdicio:.0%}"
                 + (f" · Moneda: {moneda}" if moneda else "")])
    hoja.append([])
    hoja.append(COLUMNAS)
    fila_encabezado = hoja.max_row
    for celda in hoja[fila_encabezado]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="2E7D32")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    primera = fila_encabezado + 1
    for partida in partidas:
        hoja.append(_fila(partida, desperdicio))
    ultima = hoja.max_row
    hoja.append([None] * 8 + [f"Total {moneda}".strip(), f"=SUM(J{primera}:J{ultima})"])
    hoja.cell(hoja.max_row, 9).font = Font(bold=True)
    hoja.cell(hoja.max_row, 10).font = Font(bold=True)
    for fila in hoja.iter_rows(min_row=primera, max_row=hoja.max_row):
        for columna in (3, 5):
            fila[columna - 1].number_format = "#,##0.00"
        fila[8].number_format = "#,##0.0000"
        fila[9].number_format = "#,##0.00"
    anchos = [16, 38, 11, 8, 12, 8, 12, 46, 14, 14]
    for i, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = hoja.cell(primera, 1)
    libro.save(ruta)


def resumen_por_categoria(partidas, desperdicio=0.0) -> List[tuple]:
    """[(categoría, subtotal, partidas sin precio)] en el orden de las categorías."""
    resumen = {}
    for p in partidas:
        subtotal, sin_precio = resumen.get(p.categoria, (0.0, 0))
        costo = p.costo(desperdicio)
        resumen[p.categoria] = (subtotal + (costo or 0.0), sin_precio + (costo is None))
    return [(c, *resumen[c]) for c in ORDEN_CATEGORIAS + sorted(set(resumen) - set(ORDEN_CATEGORIAS))
            if c in resumen]
